#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""selftest.py —— order-report 技能 + 报表文档规范的自检套件

不需要任何外部数据库：用 sqlite 造一份与样板文档同名的固定数据，然后

  A. 从 `samples/达人订单报表.md` 里**解析出** §7 的 SQL 模板与「输出列」表，
     真的跑一遍 —— 于是「规范里的模板可执行」是被验证的，不是声称的；
  B. 把所有安全护栏逐条撞一遍。

    python fork/report-automation/tests/selftest.py

对应自检清单的第 3、4 层（脚本本身 / 护栏）与文档一致性。第 0~2 层（沙箱出网 /
依赖 / 凭证注入）必须在真实沙箱里验，本地测不了。
"""

from __future__ import annotations

import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile

# Windows 控制台默认 GBK，中文会直接抛 UnicodeEncodeError 或被写成乱码
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
RUN_REPORT = os.path.normpath(os.path.join(
    HERE, "..", "skill", "order-report", "scripts", "run_report.py"
))
SAMPLE_DOC = os.path.normpath(os.path.join(
    HERE, "..", "samples", "达人订单报表.md"
))
TEMPLATE_DOC = os.path.normpath(os.path.join(
    HERE, "..", "templates", "report-doc-template.md"
))

# 固定数据：表名与列名刻意与样板文档 §4 保持一致，
# 订单号用 19 位（超过 2**53），用来验证 xlsx 里不会变成科学计数法
FIXTURE_SQL = """
CREATE TABLE t_daren_order (
    order_id        INTEGER,
    order_time      TEXT,
    influencer_id   TEXT,
    influencer_name TEXT,
    product_name    TEXT,
    shop_name       TEXT,
    quantity        INTEGER,
    pay_amount      REAL,
    order_status    TEXT,
    refund_status   TEXT
);
INSERT INTO t_daren_order VALUES
 (2026080100000000001, '2026-08-01 10:00:00', 'D001', '张三', '面膜A', '官方旗舰店', 2, 199.00, 'paid',    'none'),
 (2026080100000000002, '2026-08-01 11:30:00', 'D001', '张三', '面膜B', '官方旗舰店', 1,  99.50, 'paid',    'none'),
 (2026080500000000003, '2026-08-05 09:15:00', 'D002', '李四', '精华C', '海外专营店', 3, 597.00, 'paid',    'refunded'),
 (2026080600000000006, '2026-08-06 14:20:00', 'D002', '李四', '面霜F', '官方旗舰店', 2, 376.00, 'pending', 'none'),
 (2026072000000000004, '2026-07-20 16:45:00', 'D002', '李四', '面霜D', '官方旗舰店', 1, 288.00, 'paid',    'none'),
 (2026080900000000005, '2026-08-09 20:05:00', 'D003', '王五', '眼霜E', '官方旗舰店', 4, 796.00, 'paid',    'none'),
 (2026080700000000007, '2026-08-07 08:00:00', 'D004', '赵六', '测试品', '测试店铺',   9,  99.00, 'paid',    'none');
"""

PARAMS = {"start_date": "2026-08-01", "end_date": "2026-09-01", "influencer": None}

results = []


# --------------------------------------------------------------------------
# 从样板文档解析 §7 的 SQL 模板与「输出列」表
#
# 这一段本身就是「文档规范可被机器消费」的证明：智能体要做的事情与此同构 ——
# 读文档 §7，取出 SQL 与列名→中文表头映射。
# --------------------------------------------------------------------------

def parse_sample_doc(path):
    """返回 {小节标题: {"sql":..., "headers": {列名: 中文表头}, "columns": [列名]}}。"""
    with open(path, "r", encoding="utf-8") as fh:
        text = fh.read()

    sections = {}
    # 按 ### 7.1 / ### 7.2 ... 切分
    parts = re.split(r"^### (7\.\d+[^\n]*)$", text, flags=re.MULTILINE)
    for i in range(1, len(parts), 2):
        heading = parts[i].strip()
        body = parts[i + 1]

        sql_match = re.search(r"```sql\n(.*?)```", body, re.DOTALL)
        if not sql_match:
            raise AssertionError(f"样例文档的「{heading}」小节里没有 ```sql 代码块")
        sql = sql_match.group(1)

        headers = {}
        for line in body.splitlines():
            cells = [c.strip() for c in line.strip().strip("|").split("|")] if line.strip().startswith("|") else []
            if len(cells) < 2 or cells[0] in ("列名", "---") or set(cells[0]) <= set("-: "):
                continue
            headers[cells[0]] = cells[1]

        sections[heading] = {"sql": sql, "headers": headers, "columns": list(headers)}
    return sections


def split_top_level(text, sep=","):
    out, depth, current = [], 0, []
    for ch in text:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        if ch == sep and depth == 0:
            out.append("".join(current))
            current = []
            continue
        current.append(ch)
    out.append("".join(current))
    return [s.strip() for s in out if s.strip()]


def select_aliases(sql):
    """从 SQL 里取出 SELECT 列表的别名集合（用于校验文档 §7 的「列名」是否对得上）。"""
    match = re.search(r"\bSELECT\b(.*?)\bFROM\b", sql, re.DOTALL | re.IGNORECASE)
    if not match:
        raise AssertionError(f"解析不出 SELECT ... FROM 结构：{sql[:80]!r}")
    aliases = []
    for item in split_top_level(match.group(1)):
        item = " ".join(item.split())
        as_match = re.search(r"\bAS\s+([A-Za-z_][A-Za-z0-9_]*)\s*$", item, re.IGNORECASE)
        if as_match:
            aliases.append(as_match.group(1))
            continue
        tail = re.search(r"([A-Za-z_][A-Za-z0-9_]*)\s*$", item)
        if tail:
            aliases.append(tail.group(1))
    return aliases


# --------------------------------------------------------------------------
# 运行器
# --------------------------------------------------------------------------

def run(sql, params=None, extra_args=(), env_overrides=None, files=True, tmpdir=None):
    """跑一次 run_report.py，返回 (exit_code, 最后一行 JSON)。"""
    sql_path = os.path.join(tmpdir, "query.sql")
    with open(sql_path, "w", encoding="utf-8") as fh:
        fh.write(sql)

    env = {
        "PATH": os.environ.get("PATH", ""),
        "SYSTEMROOT": os.environ.get("SYSTEMROOT", ""),   # Windows 上 python 依赖它
        "PYTHONIOENCODING": "utf-8",
    }
    if files:
        env.update({
            "REPORT_DB_DRIVER": "sqlite",
            "REPORT_DB_NAME": os.path.join(tmpdir, "fixture.db"),
            "WEKNORA_SKILL_OUTPUT_DIR": os.path.join(tmpdir, "out"),
        })
    for key, value in (env_overrides or {}).items():
        if value is None:
            env.pop(key, None)
        else:
            env[key] = value

    cmd = [sys.executable, RUN_REPORT, "--sql-file", sql_path,
           "--params", json.dumps(params if params is not None else PARAMS, ensure_ascii=False)]
    cmd.extend(extra_args)

    proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", env=env)
    payload = {}
    for line in (proc.stdout or "").strip().splitlines():
        if line.strip().startswith("{"):
            payload = json.loads(line.strip())
    return proc.returncode, payload, proc


def check(name, condition, detail=""):
    results.append((name, bool(condition), detail))
    print(f"  [{'PASS' if condition else 'FAIL'}] {name}"
          + (f"  — {detail}" if detail and not condition else ""))


def write_file(tmpdir, name, content):
    path = os.path.join(tmpdir, name)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(content)
    return path


def main():
    tmpdir = tempfile.mkdtemp(prefix="order-report-selftest-")
    try:
        db_path = os.path.join(tmpdir, "fixture.db")
        conn = sqlite3.connect(db_path)
        conn.executescript(FIXTURE_SQL)
        conn.commit()
        conn.close()

        sections = parse_sample_doc(SAMPLE_DOC)
        keys = sorted(sections)
        if len(keys) < 3:
            raise AssertionError(f"样例文档应至少含 3 个 §7 小节，实际 {len(keys)} 个")

        detail = sections[keys[0]]
        by_influencer = sections[keys[1]]
        by_day = sections[keys[2]]

        merged_headers = {}
        for s in sections.values():
            merged_headers.update(s["headers"])
        headers_arg = json.dumps(merged_headers, ensure_ascii=False)

        print(f"\n[0] 文档 → 模板解析（样本：{os.path.basename(SAMPLE_DOC)}）")
        for key in keys:
            check(f"解析出「{key}」的 SQL 与输出列",
                  "SELECT" in sections[key]["sql"].upper() and sections[key]["columns"],
                  f"columns={sections[key]['columns']}")

        print("\n[0b] 文档自洽性：§7 的 SQL 别名 vs「输出列」表的列名")
        for key in keys:
            aliases = select_aliases(sections[key]["sql"])
            declared = sections[key]["columns"]
            check(f"「{key}」两侧列名一致",
                  sorted(aliases) == sorted(declared),
                  f"SQL 别名={sorted(aliases)} / 文档声明={sorted(declared)}")
        check("模板不含硬编码日期（必须用 :占位符）",
              all(not re.search(r"'20\d\d-\d\d-\d\d", s["sql"])
                  for s in sections.values()),
              "发现写死的日期字面量")
        check("模板不含 LIMIT（行数上限由执行器统一压制）",
              all(not re.search(r"\bLIMIT\b", s["sql"], re.IGNORECASE)
                  for s in sections.values()),
              "模板里出现了 LIMIT")
        check("模板用显式 JOIN，未用逗号连接表",
              all(not re.search(r"\bFROM\s+\w+\s*,", s["sql"], re.IGNORECASE)
                  for s in sections.values()),
              "出现 FROM a, b 形式，会让表名白名单漏检")

        print("\n[1] 正向链路：文档模板 → 查询 → xlsx")
        code, payload, proc = run(detail["sql"], extra_args=[
            "--out", "达人订单报表_20260801_20260831.xlsx",
            "--sheet", "明细",
            "--headers", headers_arg,
            "--summary-sql-file", write_file(tmpdir, "summary.sql", by_influencer["sql"]),
            "--max-rows", "1000",
        ], tmpdir=tmpdir)
        check("详情 + 汇总整体成功", code == 0 and payload.get("ok"),
              f"exit={code} payload={payload}")
        # 2026-08 有效行：8/1 两行 + 8/5(已退款,明细模板不过滤) + 8/6(待支付) + 8/9；
        # 8/7 是测试店铺被模板排除 → 5 行
        check("行数符合模板语义（测试店铺被排除）", payload.get("rows") == 5,
              f"rows={payload.get('rows')}")
        check("sheet 列表含 明细/汇总/查询说明",
              payload.get("sheets") == ["明细", "汇总", "查询说明"],
              f"sheets={payload.get('sheets')}")

        out_file = os.path.join(tmpdir, "out", payload.get("file", ""))
        check("xlsx 文件已生成", os.path.isfile(out_file), out_file)

        if os.path.isfile(out_file):
            from openpyxl import load_workbook
            wb = load_workbook(out_file)
            check("工作表齐全", wb.sheetnames == ["明细", "汇总", "查询说明"], f"{wb.sheetnames}")

            sheet = wb["明细"]
            check("表头用了文档「中文表头」列", sheet["A1"].value == "订单号",
                  f"A1={sheet['A1'].value!r}")
            check("首行冻结 + 自动筛选", sheet.freeze_panes == "A2" and bool(sheet.auto_filter.ref),
                  f"freeze={sheet.freeze_panes} ref={sheet.auto_filter.ref}")
            first_id = sheet["A2"].value
            check("19 位订单号按文本写入（未变科学计数法）",
                  isinstance(first_id, str) and first_id.isdigit() and len(first_id) == 19,
                  f"A2={first_id!r} type={type(first_id).__name__}")
            check("长数字列格式锁为文本", sheet["A2"].number_format == "@",
                  f"fmt={sheet['A2'].number_format!r}")

            summary = wb["汇总"]
            check("汇总表用自己那节的表头（未串到明细列名）",
                  [summary.cell(1, i).value for i in range(1, 5)] ==
                  ["达人", "订单数", "总数量", "总金额"],
                  f"{[summary.cell(1, i).value for i in range(1, 5)]}")
            # 有效订单（paid 且非 refunded，排除测试店铺）：张三 2+1=3 件、王五 4 件 → 合计 7
            # 被排除：8/5 已退款、8/6 待支付、8/7 测试店铺
            totals = [summary.cell(r, 3).value for r in range(2, summary.max_row + 1)]
            check("汇总只算有效订单（已退款/待支付/测试单被排除）",
                  sum(totals or []) == 7 and sorted(totals) == [3, 4],
                  f"各达人总数量={totals}")

            meta = wb["查询说明"]
            meta_text = "\n".join(str(c.value) for row in meta.iter_rows() for c in row if c.value)
            check("查询说明记录了参数与 SQL", "start_date" in meta_text and "SELECT" in meta_text)

        # 有效订单只落在 8/1（张三两单）与 8/9（王五）两天 → 2 行
        code, payload, _ = run(by_day["sql"], extra_args=[
            "--out", "按天.xlsx", "--headers", headers_arg], tmpdir=tmpdir)
        check("§7.3 按天模板跑通且天数正确（DATE() 三库通用）",
              code == 0 and payload.get("rows") == 2, f"exit={code} rows={payload.get('rows')}")

        print("\n[2] 注入与破坏性语句必须被拒")
        for name, sql, expected in [
            ("DROP TABLE 被拒", "DROP TABLE t_daren_order", 3),
            ("DELETE 被拒", "DELETE FROM t_daren_order", 3),
            ("UPDATE 被拒", "UPDATE t_daren_order SET quantity = 0", 3),
            ("多语句（SELECT; DELETE）被拒",
             "SELECT 1 FROM t_daren_order; DELETE FROM t_daren_order", 3),
            ("多语句（SELECT; SELECT）被拒",
             "SELECT 1 FROM t_daren_order; SELECT 2 FROM t_daren_order", 3),
            ("SELECT INTO OUTFILE 被拒",
             "SELECT * INTO OUTFILE '/tmp/x' FROM t_daren_order", 3),
            ("带锁读取 FOR UPDATE 被拒",
             "SELECT * FROM t_daren_order FOR UPDATE", 3),
            ("sleep() 被拒", "SELECT sleep(10) FROM t_daren_order", 3),
            ("非 SELECT 开头被拒", "SHOW TABLES", 3),
            ("孤儿分号被拒", "SELECT 1 FROM t_daren_order;;", 3),
        ]:
            code, payload, _ = run(sql, extra_args=["--dry-run"], tmpdir=tmpdir)
            check(name, code == expected,
                  f"exit={code}（期望 {expected}）err={payload.get('error')}")

        print("\n[3] 骨架化必须不误判（注释与字符串里的关键字不算数）")
        for name, sql in [
            ("字符串里的 'drop' 不误判",
             "SELECT order_id FROM t_daren_order WHERE shop_name = 'table-drop-shop'"),
            ("字符串里的分号不误判为多语句",
             "SELECT order_id FROM t_daren_order WHERE shop_name = 'a;b'"),
            ("-- 注释里的关键字不误判",
             "SELECT order_id FROM t_daren_order -- drop table t_daren_order\n"),
            ("/* */ 注释里的关键字不误判",
             "SELECT order_id /* drop table */ FROM t_daren_order"),
            ("首行注释不影响「以 SELECT 开头」判定",
             "-- 说明文字\nSELECT order_id FROM t_daren_order"),
            ("字符串里的 :冒号 不被当参数（时间格式）",
             "SELECT order_id FROM t_daren_order WHERE order_time > '2026-08-01 12:30'"),
            ("WITH CTE 允许",
             "WITH recent AS (SELECT * FROM t_daren_order) SELECT order_id FROM recent"),
            ("大小写与换行不影响判定",
             "select\n  order_id\nfrom t_daren_order"),
        ]:
            code, payload, _ = run(sql, params={}, extra_args=["--dry-run"], tmpdir=tmpdir)
            check(name, code == 0, f"exit={code} err={payload.get('error')}")

        print("\n[4] 参数绑定")
        code, payload, _ = run(
            "SELECT order_id FROM t_daren_order WHERE influencer_name = :influencer",
            {"influencer": "张'三"}, extra_args=["--dry-run"], tmpdir=tmpdir)
        check("单引号被转义，未形成注入",
              code == 0 and "''" in payload.get("rendered_sql", ""),
              f"sql={payload.get('rendered_sql')}")
        code, payload, _ = run(
            "SELECT order_id FROM t_daren_order WHERE influencer_name IN :names",
            {"names": ["张三", "李四"]}, extra_args=["--dry-run"], tmpdir=tmpdir)
        check("列表参数展开为 IN 列表",
              code == 0 and "('张三', '李四')" in payload.get("rendered_sql", ""),
              f"sql={payload.get('rendered_sql')}")
        code, payload, _ = run(
            "SELECT order_id FROM t_daren_order "
            "WHERE order_time >= :start_date AND order_time < :end_date",
            {"start_date": "2026-08-01"}, extra_args=["--dry-run"], tmpdir=tmpdir)
        check("缺少参数时退出码 2 并列出缺哪个",
              code == 2 and "end_date" in payload.get("error", ""),
              f"exit={code} err={payload.get('error')}")
        code, payload, _ = run(
            "SELECT order_id FROM t_daren_order WHERE influencer_name = :x",
            {"y": "张三"}, extra_args=["--dry-run"], tmpdir=tmpdir)
        check("参数名拼错时报缺参数",
              code == 2 and "x" in payload.get("error", ""),
              f"exit={code} err={payload.get('error')}")
        code, payload, _ = run("SELECT order_id FROM t_daren_order WHERE 1=1",
                               {"unused": "值"}, extra_args=["--dry-run"], tmpdir=tmpdir)
        check("多余参数只告警不报错", code == 0, f"exit={code}")

        print("\n[5] 行数上限压制")
        code, payload, _ = run("SELECT order_id FROM t_daren_order", params={},
                               extra_args=["--dry-run", "--max-rows", "100"], tmpdir=tmpdir)
        check("无 LIMIT 时外包一层 LIMIT",
              code == 0 and "AS _weknora_limited LIMIT 100" in payload.get("rendered_sql", ""),
              f"sql={payload.get('rendered_sql')}")
        code, payload, _ = run("SELECT order_id FROM t_daren_order LIMIT 9999999", params={},
                               extra_args=["--dry-run", "--max-rows", "100"], tmpdir=tmpdir)
        check("过大 LIMIT 被压到上限",
              code == 0 and "LIMIT 100" in payload.get("rendered_sql", "")
              and "9999999" not in payload.get("rendered_sql", ""),
              f"sql={payload.get('rendered_sql')}")
        code, payload, _ = run("SELECT order_id FROM t_daren_order LIMIT 5", params={},
                               extra_args=["--dry-run", "--max-rows", "100"], tmpdir=tmpdir)
        check("较小 LIMIT 被保留",
              code == 0 and "LIMIT 5" in payload.get("rendered_sql", ""),
              f"sql={payload.get('rendered_sql')}")

        print("\n[6] 表名白名单")
        code, payload, _ = run(detail["sql"],
                               extra_args=["--dry-run", "--allowed-tables", "t_other"], tmpdir=tmpdir)
        check("白名单外的表被拒", code == 3 and "白名单" in payload.get("error", ""),
              f"exit={code} err={payload.get('error')}")
        code, payload, _ = run(detail["sql"],
                               extra_args=["--dry-run", "--allowed-tables", "t_daren_order"], tmpdir=tmpdir)
        check("白名单内的表通过", code == 0, f"exit={code} err={payload.get('error')}")

        print("\n[7] 空结果与配置缺失")
        code, payload, _ = run(
            "SELECT order_id FROM t_daren_order WHERE order_time >= '2030-01-01'",
            params={}, extra_args=["--out", "empty.xlsx"], tmpdir=tmpdir)
        check("空结果集正常出文件且 rows=0",
              code == 0 and payload.get("rows") == 0, f"exit={code} rows={payload.get('rows')}")
        code, payload, _ = run(detail["sql"], files=False, tmpdir=tmpdir)
        check("缺少数据库环境变量时退出码 2 且提示去哪填",
              code == 2 and "沙箱密钥" in payload.get("error", ""),
              f"exit={code} err={payload.get('error')}")
        code, payload, _ = run(detail["sql"], files=False,
                               env_overrides={"REPORT_DB_HOST": "h", "REPORT_DB_USER": "u",
                                              "REPORT_DB_NAME": "d", "REPORT_DB_PASSWORD": None},
                               tmpdir=tmpdir)
        check("只缺密码时准确点名 REPORT_DB_PASSWORD",
              code == 2 and "REPORT_DB_PASSWORD" in payload.get("error", ""),
              f"exit={code} err={payload.get('error')}")
        code, payload, _ = run(detail["sql"], params={},
                               extra_args=["--sql-file", os.path.join(tmpdir, "nope.sql")],
                               tmpdir=tmpdir)
        check("sql 文件不存在时给出可读错误", code == 2, f"exit={code} err={payload.get('error')}")

        print("\n[8] 文件名清洗")
        code, payload, _ = run(detail["sql"],
                               extra_args=["--out", "达人 订单(报表).xlsx"], tmpdir=tmpdir)
        check("文件名里的空格与括号被清洗掉",
              code == 0 and " " not in payload.get("file", "") and "(" not in payload.get("file", ""),
              f"file={payload.get('file')}")

        print("\n[9] 空模板与样板的结构一致性")
        with open(TEMPLATE_DOC, encoding="utf-8") as fh:
            template = fh.read()
        sample_heads = re.findall(r"^## (\d+\..+)$", open(SAMPLE_DOC, encoding="utf-8").read(), re.M)
        tpl_heads = re.findall(r"^## (\d+\..+)$", template, re.M)
        check("空模板与样板章节标题完全一致（八节顺序固定）",
              sample_heads == tpl_heads,
              f"样板={sample_heads} / 模板={tpl_heads}")

    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

    passed = sum(1 for _, ok, _ in results if ok)
    total = len(results)
    print(f"\n{'=' * 64}\n自检结果：{passed}/{total} 通过")
    if passed != total:
        print("\n失败项：")
        for name, ok, detail in results:
            if not ok:
                print(f"  - {name}  {detail}")
        return 1
    print("全部通过。\n"
          "已验证：文档规范可被机器解析、样板模板真的可执行、脚本链路与全部安全护栏、"
          "文档自洽性。\n"
          "未验证（需在真实沙箱里做）：沙箱出网、技能依赖安装、凭证注入 —— 见 checklist.md 第 0~2 层。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
